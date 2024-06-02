import model.Banco.BD as conexao
import openpyxl

#Login
def LoginUsuario(login, senha):
    conexaoBD = conexao.iniciaConexao()
    query = "SELECT * FROM usuarios WHERE login = %s and senha = %s;"
    parametros = (login, senha) 

    cursorBD = conexaoBD.cursor()
    cursorBD.execute(query, parametros)
    listaResultado = cursorBD.fetchone()
    cursorBD.close()
    conexaoBD.close()
    return listaResultado


#Cadastrar Novo Ususario

def CadastrarUsuario(login, senha):
    conexaoBD = conexao.iniciaConexao()
    query = "SELECT * FROM usuarios WHERE login = %s and senha = %s;"
    parametros = (login, senha) 

    cursorBD = conexaoBD.cursor()
    cursorBD.execute(query, parametros)
    listaResultado = cursorBD.fetchone()

    if(listaResultado != None):
        cadastrar = "INSERT INTO usuarios (login, senha) VALUES (%s, %s);"
        cursorBD.execute(cadastrar, parametros)
        conexaoBD.commit()
    else:
         return False

    cursorBD.close()
    conexaoBD.close()

#Cadastrar Turma
def CadastrarTurma(planilha):
    book = openpyxl.load_workbook(planilha)
    ws = book.active
    sheet = book['3º TDS "A"']  
    ws.tables

    for table in ws.tables.values():
        print(table)

    num_min = "10"
    num_max = "10"
    letra_min = "A"
    letra_max = "L"
    min = letra_min + num_min
    max = letra_max + num_max
    cell_range = sheet[min+':'+ max]
    alunos = []
    for row in cell_range:
         for cell in row:
            alunos.append(cell.value)

    if (alunos != None):
        conexaoBD = conexao.iniciaConexao()
        query = "INSERT INTO bdalunos (instituicao, nome_completo, cpf, rg, orgao_expedidor, municipio, nacionalidade, data_nascimento, curso, data_conclusão, nome_pai, nome_mae) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);"
        parametros = (alunos[0], alunos[1], alunos[2],alunos[3], alunos[4], alunos[5], alunos[6], alunos[7], str(alunos[8]), str(alunos[9]), alunos[10], alunos[11])

        cursorBD = conexaoBD.cursor()
        cursorBD.execute(query, parametros)
        conexaoBD.commit()
        
        cursorBD.close()
        conexaoBD.close()
        return alunos
    else:
        return False